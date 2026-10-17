"""Attachment contracts: content bounds, ownership, ingest, and run handoff."""
import asyncio
import io

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from mozhi_service.api import create_app
from mozhi_service.attachments import MAX_BYTES, extract_text, image_mime
from mozhi_service.runtime import Runtime
from mozhi_service.schemas import RunInput
from mozhi_service.settings import Settings
from mozhi_service.store import Store


def setup(tmp_path):
    settings = Settings(db_path=str(tmp_path / "data.db"), model_key="fake-secret")
    store = Store(settings.db_path)
    first = store.create_user({"nickname": "Alice", "daily_run_limit": 10, "enabled_tools": []})
    second = store.create_user({"nickname": "Bob", "daily_run_limit": 10, "enabled_tools": []})
    return settings, store, TestClient(create_app(settings, store)), first, second


def token(user):
    return {"Authorization": "Bearer " + user["token"]}


def docx_bytes():
    document = Document()
    document.add_paragraph("甲方合同的付款期限是三十天")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "成绩"
    table.cell(1, 0).text = "学生甲"
    table.cell(1, 1).text = "90"
    data = io.BytesIO()
    document.save(data)
    return data.getvalue()


def xlsx_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销售"
    sheet.append(["日期", "毛利"])
    sheet.append(["2026-10-10", 1800])
    data = io.BytesIO()
    workbook.save(data)
    return data.getvalue()


def test_extract_word_and_excel():
    word = extract_text("合同.docx", docx_bytes())
    excel = extract_text("数据.xlsx", xlsx_bytes())
    assert "付款期限" in word and "学生甲 | 90" in word
    assert "工作表: 销售" in excel and "毛利" in excel and "1800" in excel


def test_upload_auth_size_format_and_user_isolation(tmp_path):
    settings, store, client, alice, bob = setup(tmp_path)
    body = docx_bytes()
    no_auth = client.post("/v1/agent/attachments", files={"file": ("x.docx", body)})
    assert no_auth.status_code == 401
    upload = client.post("/v1/agent/attachments", headers=token(alice),
                         files={"file": ("合同.docx", body)})
    assert upload.status_code == 200, upload.text
    attachment = upload.json()
    assert attachment["name"] == "合同.docx"
    assert "付款期限" in store.attachment(alice["user"]["id"], attachment["id"])["content"]
    forbidden = RunInput(client_request_id="attachment-test-bob-01",
                         messages=[{"role": "user", "content": "分析合同"}],
                         attachments=[attachment["id"]]).model_dump()
    response = client.post("/v1/agent/runs", headers=token(bob), json=forbidden)
    assert response.status_code == 404
    unsupported = client.post("/v1/agent/attachments", headers=token(alice),
                              files={"file": ("script.py", b"print('hello')")})
    assert unsupported.status_code == 415
    too_big = client.post("/v1/agent/attachments", headers=token(alice),
                          files={"file": ("fake.png", b"x" * (MAX_BYTES + 1))})
    assert too_big.status_code == 413


def test_excel_upload_and_agent_context_handoff(tmp_path):
    settings, store, client, alice, _ = setup(tmp_path)
    upload = client.post("/v1/agent/attachments", headers=token(alice),
                         files={"file": ("数据.xlsx", xlsx_bytes())})
    assert upload.status_code == 200, upload.text
    item = upload.json()
    request = RunInput(client_request_id="attachment-request-001",
                       messages=[{"role": "user", "content": "分析数据"}],
                       attachments=[item["id"]]).model_dump()
    run = client.post("/v1/agent/runs", headers=token(alice), json=request)
    assert run.status_code == 200
    assert store.claim() == run.json()["id"]
    class FakeModel:
        def __init__(self):
            self.messages = None
        async def complete(self, messages, tools, model, final_only):
            import copy
            self.messages = copy.deepcopy(messages)
            return {"content": "毛利 1800"}, {"prompt_tokens": 3, "completion_tokens": 4}
    model = FakeModel()
    asyncio.run(Runtime(store, settings, model).execute(run.json()["id"]))
    assert store.run(run.json()["id"])["state"] == "completed"
    assert "1800" in model.messages[-1]["content"]
    assert "不可信来源" in model.messages[-1]["content"]


def test_image_upload_routes_to_vision_and_only_keeps_analysis(tmp_path, monkeypatch):
    settings, store, client, alice, _ = setup(tmp_path)
    called = []
    async def fake_vision(settings, name, blob):
        called.append(name)
        return "图片内容：异常折线图，2026年10月毛利降低"
    monkeypatch.setattr("mozhi_service.api.recognize_image", fake_vision)
    response = client.post("/v1/agent/attachments", headers=token(alice),
                           files={"file": ("chart.png", b"\x89PNG\r\n\x1a\nexample")})
    assert response.status_code == 200
    item = response.json()
    saved = store.attachment(alice["user"]["id"], item["id"])
    assert item["kind"] == "image" and called == ["chart.png"]
    assert "异常折线图" in saved["content"]
    assert not saved["content"].startswith("data:image")


def test_vision_signature_and_schema_guardrails():
    from mozhi_service.store import Problem
    try:
        image_mime("fake.jpg", b"not-a-jpeg")
        assert False, "must reject masquerading image"
    except Problem as error:
        assert error.status == 422
    from pydantic import ValidationError
    try:
        RunInput(client_request_id="attachment-schema-001", messages=[{"role": "user", "content": "hi"}],
                 attachments=["bad-id"])
        assert False
    except ValidationError:
        pass

def test_upload_attempt_rate_limit(tmp_path):
    _, store, client, alice, _ = setup(tmp_path)
    doc = docx_bytes()
    for _ in range(6):
        r = client.post("/v1/agent/attachments", headers=token(alice),
                        files={"file": ("test.docx", doc)})
        assert r.status_code == 200, r.text
    blocked = client.post("/v1/agent/attachments", headers=token(alice),
                          files={"file": ("test.docx", doc)})
    assert blocked.status_code == 429
